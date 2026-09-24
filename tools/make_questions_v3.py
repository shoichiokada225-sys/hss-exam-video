#!/usr/bin/env python3
"""第3版(2026-09-24): エージェント作問→検証済みの問題を tools/questions.xlsx の「本試験」シートに追記する。

入力:  raw/v2/verified/*.json（検証チームが keep/fix にした問題。drop は入れない）
       raw/v2/manifest.json（素材ID → 元ファイルの場所）
出力:  tools/questions.xlsx に行を追記（既存30問は一切変えない）
       tools/clips_v2.json（切り出し台帳。python3 tools/cut_clips_v2.py で videos/ を作る）

- 新しい問題は placeholder=1（未監修）で入る。社長が review.html で OK にすると 0 になる
- id は q31 から連番。動画ファイル名も同じ id（1問1本。同じ場面の別の問いも別ファイルにして素材台帳で追える）
- 追加列: level(初級/中級/上級) / source(根拠) / scene(何が写っているか) / media(素材ID@秒) / check_note(【社長確認】・作問/検証メモ)
- 何度実行しても同じ結果になる（media+問題文が同じ行は追記しない）

使い方:  python3 tools/make_questions_v3.py [--dry-run]
"""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_from_xlsx import CATEGORIES, HEAD, XLSX  # noqa: E402

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except (AttributeError, ValueError):
    pass

ROOT = Path(__file__).resolve().parent.parent
V2 = ROOT / "raw" / "v2"
LEDGER = ROOT / "tools" / "clips_v2.json"
EXTRA = ["level", "source", "scene", "media", "check_note", "group"]
LEVELS = ("初級", "中級", "上級")
HOME = str(Path.home())


def scene_group(media):
    """同じ場面の問いをまとめる鍵。中身が同じ別IDの素材(md5一致)は代表IDにそろえる"""
    base, _, at = media.partition("@")
    dups = V2 / "work" / "lead" / "md5_dups.json"
    if dups.exists():
        for g in json.loads(dups.read_text(encoding="utf-8")):
            if base in g:
                base = sorted(g)[0]
    return base + (f"@{at}" if at else "")


def check_note(it):
    """監修者向けメモ: 【社長確認】があれば先頭に。作問者 note と検証 note をつなぐ"""
    parts = [x for x in (it.get("review_note"), it.get("note")) if x]
    parts.sort(key=lambda x: "社長確認" not in x)
    return " / ".join(parts)


def load_items():
    man = {
        m["id"]: m
        for m in json.loads((V2 / "manifest.json").read_text(encoding="utf-8"))
    }
    items = []
    for f in sorted((V2 / "verified").glob("*.json")):
        for it in json.loads(f.read_text(encoding="utf-8")):
            if it.get("verdict", "keep") == "drop":
                continue
            m = man.get(it["media"])
            if not m:
                raise SystemExit(f"{f.name}: manifest に無い素材 {it['media']}")
            if it["category"] not in CATEGORIES:
                raise SystemExit(
                    f"{f.name}: category が不正 {it['category']!r} ({it['media']})"
                )
            if it["level"] not in LEVELS:
                raise SystemExit(
                    f"{f.name}: level が不正 {it['level']!r} ({it['media']})"
                )
            for lang in ("ja", "en", "vi", "id"):
                o = it[lang]["opts"]
                if (
                    len(o) != 4
                    or not all(str(x).strip() for x in o)
                    or not str(it[lang]["q"]).strip()
                ):
                    raise SystemExit(f"{f.name}: {it['media']} の {lang} が不完全")
                if len(set(o)) != 4:
                    raise SystemExit(
                        f"{f.name}: {it['media']} の {lang} に同じ選択肢がある"
                    )
            it["_m"] = m
            it["_file"] = f.stem
            items.append(it)
    return items


def main():
    dry = "--dry-run" in sys.argv
    items = load_items()
    wb = load_workbook(XLSX)
    ws = wb["本試験"]
    head = [c.value for c in ws[1]]
    for h in EXTRA:
        if h not in head:
            ws.cell(row=1, column=len(head) + 1, value=h)
            head.append(h)
    col = {h: i for i, h in enumerate(head)}
    existing, max_n = set(), 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        qid = row[col["id"]]
        if qid and str(qid).startswith("q") and str(qid)[1:].isdigit():
            max_n = max(max_n, int(str(qid)[1:]))
        media = row[col["media"]] if col["media"] < len(row) else None
        existing.add((media, row[col["question_ja"]]))

    ledger = (
        json.loads(LEDGER.read_text(encoding="utf-8"))
        if LEDGER.exists()
        else {"_doc": "", "clips": []}
    )
    ledger["_doc"] = (
        "第3版の切り出し台帳(make_questions_v3.py が生成)。src=元素材(~からの相対)。"
        "kind=video は start 秒から5秒、image は5秒ズーム。python3 tools/cut_clips_v2.py で videos/ に生成"
    )
    added = 0
    for it in items:
        m = it["_m"]
        media = (
            f"{it['media']}@{it['start']:g}"
            if m["kind"] == "video" and it.get("start") is not None
            else it["media"]
        )
        if (media, it["ja"]["q"]) in existing:
            continue
        max_n += 1
        qid = f"q{max_n}"
        vals = {
            "id": qid,
            "video": f"videos/{qid}.mp4",
            "category": it["category"],
            "required": 0,
            "placeholder": 1,
            "question_ja": it["ja"]["q"],
            "answer": 1,
            "explanation_ja": it["explanation_ja"],
            "level": it["level"],
            "source": it.get("source", ""),
            "scene": it.get("scene", ""),
            "media": media,
                "check_note": check_note(it),
                "group": scene_group(media),
        }
        for i in range(4):
            vals[f"opt{i + 1}_ja"] = it["ja"]["opts"][i]
        for lang in ("en", "vi", "id"):
            vals[f"question_{lang}"] = it[lang]["q"]
            for i in range(4):
                vals[f"opt{i + 1}_{lang}"] = it[lang]["opts"][i]
        r = ws.max_row + 1
        for h, v in vals.items():
            ws.cell(row=r, column=col[h] + 1, value=v)
        src = m["path"].replace(HOME + "/", "")
        ledger["clips"].append(
            {
                "video": qid,
                "media": it["media"],
                "kind": m["kind"],
                "src": src,
                "start": it.get("start") if m["kind"] == "video" else None,
                "focus": it.get("focus"), "rotate": it.get("rotate") or 0,
                "scene": it.get("scene", ""),
                "team": it["_file"],
            }
        )
        existing.add((media, it["ja"]["q"]))
        added += 1
    print(
        f"検証済み {len(items)}問 → 追記 {added}問（{f'q{max_n - added + 1}〜q{max_n}' if added else 'なし'}）"
    )
    if dry:
        print("(--dry-run: 保存しません)")
        return
    wb.save(XLSX)
    LEDGER.write_text(
        json.dumps(ledger, ensure_ascii=False, indent=1) + "\n", encoding="utf-8"
    )
    print(f"保存: {XLSX.relative_to(ROOT)} / {LEDGER.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
