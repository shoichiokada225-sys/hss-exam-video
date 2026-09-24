#!/usr/bin/env python3
# 正本 tools/questions.xlsx の未監修問 → Googleドライブの確認シート（xlsx）を作る
# 動画のドライブIDは DriveFS の metadata_sqlite_db（コピー）から引く
import shutil, sqlite3, glob, sys
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
HOME = Path.home()
DRIVE = HOME / "Library/CloudStorage/GoogleDrive-shoichi.okada225@gmail.com/マイドライブ/HSS動画試験_監修_2026-09"
FOLDER_ID = "1ozvoVyLScvBn4nNXwqedIbYfA3LoK4KC"
tmp = Path("/private/tmp/claude-501/-Users-okadashoichi/92e9fdfc-5c65-41cb-9074-61513f326210/scratchpad/dfs"); tmp.mkdir(parents=True, exist_ok=True)
for f in glob.glob(str(HOME / "Library/Application Support/Google/DriveFS/*/metadata_sqlite_db*")):
    shutil.copy(f, tmp)
db = sqlite3.connect(tmp / "metadata_sqlite_db")
ids = dict(db.execute("select i.local_title,i.id from items i join stable_parents p on p.item_stable_id=i.stable_id "
                      "join items f on f.stable_id=p.parent_stable_id where f.id=? and i.trashed=0", (FOLDER_ID,)).fetchall())
src = load_workbook(HOME / "hss-exam-video/tools/questions.xlsx")["本試験"]
h = {c.value: i for i, c in enumerate(src[1])}
rows = []
for r in src.iter_rows(min_row=2, values_only=True):
    g = lambda k: "" if k not in h or r[h[k]] is None else str(r[h[k]])
    if g("placeholder") not in ("1", "1.0") or g("review_status") == "del":
        continue
    ans = int(float(g("answer"))); opts = [g(f"opt{i}_ja") for i in range(1, 5)]
    wrong = [o for i, o in enumerate(opts, 1) if i != ans]
    note = g("check_note")
    ask = note.split("社長確認】", 1)[1] if "社長確認】" in note and "解決済" not in note else ""
    fid = ids.get(g("id") + ".mp4", "")
    if not fid or fid.startswith("local-"):
        sys.exit(f"動画のドライブIDがまだありません: {g('id')}（アップロード待ち）")
    link = f'=HYPERLINK("https://drive.google.com/file/d/{fid}/view","▶ 再生")'
    tr = lambda l: g(f"question_{l}") + " → " + g(f"opt{ans}_{l}")
    rows.append([g("id"), link, g("level"), g("category"), g("question_ja"), opts[ans - 1], *wrong, "", "", "",
                 ask, g("explanation_ja"), g("scene"), g("source"), tr("en"), tr("vi"), tr("id")])
HEAD = ["問題ID", "動画", "難易度", "カテゴリ", "問題", "正解", "誤答1", "誤答2", "誤答3",
        "判定（OK/修正/削除）", "メモ（直す内容・理由）", "確認者", "【要判断】の論点", "解説",
        "写っているもの", "根拠（swine-manual）", "英語（問題 → 正解）", "ベトナム語（問題 → 正解）", "インドネシア語（問題 → 正解）"]
wb = Workbook(); gd = wb.active; gd.title = "使い方"
guide = [f"HSS動画試験 追加問題（{len(rows)}問）の確認シート", "",
 "1. 「問題」シートの各行の「▶ 再生」を押すと、5秒の動画が開きます（ブラウザの新しいタブ）。",
 "2. 動画を見て、問題・正解・誤答が正しいかを確認してください。",
 "3. 「判定」欄のプルダウンで OK / 修正 / 削除 を選び、修正・削除のときは「メモ」に理由や直し方を書いてください。",
 "4. 「確認者」欄にお名前を書いてください（複数人で分担できます）。", "",
 "・正解は「正解」列です。本番の試験では選択肢の順番はランダムに並び替わります。",
 "・本番では毎回、初級12・中級12・上級6の配分で30問が出ます。同じ場面の問いは1回に1問だけです。",
 "・難易度（初級/中級/上級）やカテゴリで絞り込みたいときは、見出し行のフィルタを使ってください。",
 "・英語/ベトナム語/インドネシア語の列は、翻訳の確認ができる方だけ見てください。",
 "・この動画と問題はまだ公開前です。外部に転送しないでください。"]
for i, t in enumerate(guide, 1):
    c = gd.cell(row=i, column=1, value=t)
    if i == 1: c.font = Font(bold=True, size=14)
gd.column_dimensions["A"].width = 110
ws = wb.create_sheet("問題"); ws.append(HEAD)
for r in rows: ws.append(r)
hdr = PatternFill("solid", fgColor="DDE8F5"); inp = PatternFill("solid", fgColor="FFF7D6"); askf = PatternFill("solid", fgColor="FDE2E2")
for c in ws[1]: c.font = Font(bold=True); c.fill = hdr; c.alignment = Alignment(wrap_text=True, vertical="center")
for k, v in {"A":8,"B":9,"C":7,"D":13,"E":42,"F":26,"G":20,"H":20,"I":20,"J":12,"K":30,"L":10,"M":40,"N":44,"O":36,"P":28,"Q":40,"R":40,"S":40}.items():
    ws.column_dimensions[k].width = v
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    row[1].font = Font(color="1155CC", underline="single", bold=True); row[5].font = Font(bold=True, color="1F7A45")
    for i in (9, 10, 11): row[i].fill = inp
    if row[12].value: row[12].fill = askf
dv = DataValidation(type="list", formula1='"OK,修正,削除"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"J2:J{ws.max_row}")
ws.freeze_panes = "C2"; ws.auto_filter.ref = f"A1:S{ws.max_row}"; wb.active = 1
dst = DRIVE / sys.argv[1] if len(sys.argv) > 1 else DRIVE / "HSS動画試験_追加問題_確認シート.xlsx"
wb.save(dst); print("saved", dst.name, len(rows), "問 / 要判断", sum(1 for r in rows if r[12]))
